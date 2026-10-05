"""数据格式转换模块：CSV / TSV / JSON / YAML / XML / TOML / XLSX 互通，
并支持表格导出为 Markdown / HTML / PDF。"""
from __future__ import annotations

import csv
import json
from io import StringIO
from pathlib import Path

from .documents import html_to_pdf
from .utils import ConvertError, read_text_smart, safe_output_path, html_escape

INPUT_EXTS = {"csv", "tsv", "json", "yaml", "yml", "xml", "toml", "xlsx"}
TARGETS = [
    ("csv", "CSV 表格"), ("tsv", "TSV 表格"), ("xlsx", "Excel 表格"),
    ("json", "JSON 数据"), ("yaml", "YAML 数据"), ("xml", "XML 数据"),
    ("toml", "TOML 数据"), ("md", "Markdown 表格"), ("html", "HTML 表格"),
    ("pdf", "PDF 表格"),
]
_TABULAR = {"csv", "tsv", "xlsx"}


def can_convert(src_ext: str, dst_ext: str) -> bool:
    return src_ext in INPUT_EXTS and dst_ext in {t for t, _ in TARGETS}


def convert(src: Path, target: str, out_dir: Path | None, options: dict) -> list[Path]:
    overwrite = options.get("overwrite", False)
    dst = safe_output_path(src, out_dir, target, overwrite)
    src_ext = src.suffix.lower().lstrip(".")

    if target in ("md", "html", "pdf"):
        records = load_records(src)
        if not records:
            raise ConvertError("数据为空或不是“表头+行”结构，无法生成表格")
        headers = list(records[0].keys())
        rows = [[_cell(r.get(h)) for r in records] for h in headers]
        rows = list(map(list, zip(*rows))) if rows else []
        if target == "md":
            _to_md_table(headers, rows, dst)
        elif target == "html":
            _to_html_table(headers, rows, dst, src.stem)
        else:
            _to_pdf_table(headers, rows, dst, src.stem)
        return [dst]

    if target in _TABULAR or target == "csv":
        records = load_records(src)
        if not records:
            raise ConvertError("数据不是“表头+多行记录”结构，无法转为表格；请转为 JSON/YAML")
        if target == "csv":
            _to_csv(records, dst, ",")
        elif target == "tsv":
            _to_csv(records, dst, "\t")
        else:
            _to_xlsx(records, dst)
        return [dst]

    # 嵌套结构目标
    data = load_any(src)
    if target == "json":
        dst.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str),
                       encoding="utf-8")
    elif target in ("yaml", "yml"):
        import yaml
        dst.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False),
                       encoding="utf-8")
    elif target == "xml":
        import xmltodict
        # xmltodict 的 unparse 要求顶层恰好一个根节点：
        # 列表包一层字典；多键字典再包一层统一根节点
        if isinstance(data, dict):
            payload = data if len(data) == 1 else {"root": data}
        elif isinstance(data, list) and data and all(isinstance(i, dict) for i in data):
            payload = {"records": {"record": data}}
        else:
            payload = {"items": {"item": data}}
        dst.write_text(xmltodict.unparse(payload, pretty=True), encoding="utf-8")
    elif target == "toml":
        import tomli_w
        payload = data if isinstance(data, dict) else {"records": data}
        try:
            dst.write_bytes(tomli_w.dumps(payload).encode("utf-8"))
        except Exception as e:
            raise ConvertError(f"该数据结构不适合 TOML（{e}）；建议转为 JSON/YAML") from e
    else:  # pragma: no cover
        raise ConvertError(f"数据模块不支持转换为 {target}")
    return [dst]


# ---------------------------------------------------------------- 读取

def load_any(src: Path):
    src_ext = src.suffix.lower().lstrip(".")
    if src_ext in ("yaml", "yml"):
        import yaml
        return yaml.safe_load(read_text_smart(src))
    if src_ext == "json":
        return json.loads(read_text_smart(src))
    if src_ext == "xml":
        import xmltodict
        data = xmltodict.parse(read_text_smart(src))
        if isinstance(data, dict) and len(data) == 1:
            return next(iter(data.values()))
        return data
    if src_ext == "toml":
        import tomllib
        with open(src, "rb") as f:
            return tomllib.load(f)
    records = load_records(src)
    return records


def load_records(src: Path) -> list[dict]:
    """把数据源解析为 [{列: 值}, …]；嵌套对象会被压平为字符串。"""
    src_ext = src.suffix.lower().lstrip(".")
    if src_ext == "csv":
        return _csv_records(src, ",")
    if src_ext == "tsv":
        return _csv_records(src, "\t")
    if src_ext == "xlsx":
        return _xlsx_records(src)
    if src_ext in ("json", "yaml", "yml", "xml", "toml"):
        data = load_any(src)
        records = _find_records(data)
        return [_flatten(r) for r in records]
    return []


def _csv_records(src: Path, delimiter: str) -> list[dict]:
    text = read_text_smart(src)
    if delimiter == ",":
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            delimiter = dialect.delimiter
        except csv.Error:
            pass
    reader = csv.DictReader(StringIO(text), delimiter=delimiter)
    return [{(k or "列"): (_clean(v)) for k, v in row.items()} for row in reader]


def _xlsx_records(src: Path) -> list[dict]:
    head = src.read_bytes()[:512]
    if head[:4] == b"\xd0\xcf\x11\xe0":
        raise ConvertError("该文件实际是旧版 Excel .xls 格式（扩展名与内容不符），无法直接读取。"
                           "请用 Excel/WPS 打开后「另存为 .xlsx」再转换。")
    if head[:4] != b"PK\x03\x04":
        # 网上系统导出的“Excel”常是网页表格伪装的，直接按 HTML 表格解析
        return _html_table_records(src)
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ConvertError("缺少 openpyxl 依赖") from e
    wb = load_workbook(src, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h) if h is not None else f"列{i + 1}" for i, h in enumerate(rows[0])]
    records = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        records.append({h: _clean(v) for h, v in zip(headers, row)})
    return records


def _html_table_records(src: Path) -> list[dict]:
    """从伪装成 .xlsx 的 HTML 文件中解析第一个表格为记录列表。"""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(read_text_smart(src), "html.parser")
    table = soup.find("table")
    if table is None:
        raise ConvertError("该文件既不是有效的 Excel，也不含网页表格，无法解析")
    rows = table.find_all("tr")
    if not rows:
        return []
    headers = [c.get_text(" ", strip=True) or f"列{i + 1}"
               for i, c in enumerate(rows[0].find_all(["th", "td"]))]
    records = []
    for row in rows[1:]:
        cells = [c.get_text(" ", strip=True) for c in row.find_all(["th", "td"])]
        if not any(cells):
            continue
        records.append({h: _clean(v) for h, v in zip(headers, cells)})
    return records


def _find_records(data) -> list:
    """从任意结构中找出记录列表；找不到则抛错提示。"""
    if isinstance(data, list) and data and all(isinstance(i, dict) for i in data):
        return data
    if isinstance(data, dict):
        for v in data.values():
            if isinstance(v, list) and v and all(isinstance(i, dict) for i in v):
                return v
        return [data]
    raise ConvertError("数据不是“对象列表”结构，无法转为表格")


def _flatten(obj, prefix: str = "") -> dict:
    out: dict = {}
    if isinstance(obj, dict):
        for k, v in obj.items():
            key = f"{prefix}.{k}" if prefix else str(k)
            if isinstance(v, dict):
                out.update(_flatten(v, key))
            elif isinstance(v, list):
                out[key] = json.dumps(v, ensure_ascii=False)
            else:
                out[key] = v
    else:
        out[prefix or "值"] = obj
    return out


def _clean(v):
    if v is None:
        return ""
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return v


def _cell(v) -> str:
    return "" if v is None else str(v)


# ---------------------------------------------------------------- 写出

def _to_csv(records: list[dict], dst: Path, delimiter: str) -> None:
    headers: list[str] = []
    for r in records:
        for k in r:
            if k not in headers:
                headers.append(k)
    with open(dst, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=headers, delimiter=delimiter, extrasaction="ignore")
        w.writeheader()
        w.writerows(records)


def _to_xlsx(records: list[dict], dst: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    headers: list[str] = []
    for r in records:
        for k in r:
            if k not in headers:
                headers.append(k)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in records:
        ws.append([r.get(h) for h in headers])
    for i, h in enumerate(headers, 1):
        width = max(10, min(50, max((len(str(r.get(h, ""))) for r in records[:200]), default=10) * 2))
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    wb.save(dst)


def _to_md_table(headers: list[str], rows: list[list[str]], dst: Path) -> None:
    def esc(s: str) -> str:
        return s.replace("|", "\\|").replace("\n", " ")
    lines = ["| " + " | ".join(esc(h) for h in headers) + " |",
             "| " + " | ".join("---" for _ in headers) + " |"]
    for row in rows:
        lines.append("| " + " | ".join(esc(c) for c in row) + " |")
    dst.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _to_html_table(headers: list[str], rows: list[list[str]], dst: Path, title: str) -> None:
    th = "".join(f"<th>{html_escape(h)}</th>" for h in headers)
    body = "".join("<tr>" + "".join(f"<td>{html_escape(c)}</td>" for c in row) + "</tr>" for row in rows)
    html = ("<style>table{border-collapse:collapse;font-family:'Microsoft YaHei',sans-serif}"
            "th,td{border:1px solid #999;padding:6px 10px}th{background:#eef}</style>\n"
            f"<h2>{html_escape(title)}</h2>\n<table>{th}{body}</table>\n")
    dst.write_text(html, encoding="utf-8")


def _to_pdf_table(headers: list[str], rows: list[list[str]], dst: Path, title: str) -> None:
    th = "".join(f"<th>{html_escape(h)}</th>" for h in headers)
    body = "".join("<tr>" + "".join(f"<td>{html_escape(c)}</td>" for c in row) + "</tr>" for row in rows)
    html = (f"<h2>{html_escape(title)}</h2>"
            f"<table border=\"1\" cellpadding=\"4\" cellspacing=\"0\" width=\"100%\">{th}{body}</table>")
    html_to_pdf(html, dst)
