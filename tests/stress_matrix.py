"""大规模转换压测：
1. 生成约 200 个内容各异的源文件（图片/文档/数据/音视频，多种变体）
2. 枚举每个文件的所有合法目标格式，逐个真实转换（并行执行）
3. 逐个校验输出：存在性、魔数、格式、内容标记
4. 报告成功率、失败清单、组合覆盖率（每种组合都必须测到）

运行:  python tests/stress_matrix.py
"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from PIL import Image, ImageDraw  # noqa: E402

from core import data_formats, documents, images, media, registry  # noqa: E402
from core.utils import ConvertError  # noqa: E402

MODULES = [("images", images), ("documents", documents),
           ("data", data_formats), ("media", media)]
PARALLEL = 6
_progress_lock = threading.Lock()
_progress_done = 0


# ================================================================ 生成源文件

def _draw_variant(img: Image.Image, k: int, marker: str) -> Image.Image:
    if img.mode == "P":
        rgb = img.convert("RGB")
        _draw_variant(rgb, k, marker)
        return rgb.quantize(colors=64)
    d = ImageDraw.Draw(img)
    if img.mode == "L":
        outline, fill, text = 40, 200, 255
    else:
        outline, fill, text = (30, 30, 30), (60 + 20 * k % 200, 90, 200), (255, 255, 255)
    d.rectangle([0, 0, img.width - 1, img.height - 1], outline=outline)
    d.ellipse([4, 4, img.width - 5, img.height - 5], fill=fill)
    d.text((6, img.height // 2 - 4), marker, fill=text)
    return img


def gen_image(path: Path, ext: str, k: int, marker: str) -> None:
    size = 48 + 13 * (k % 9)
    mode = ["RGB", "RGBA", "RGB", "L", "RGBA", "P", "RGB", "RGBA"][k % 8]
    if mode == "L":
        color = 120 + (k % 4) * 20          # 灰度
    elif mode == "P":
        color = (k % 4) + 1                  # 调色板索引
    else:
        color = [(220, 60, 60), (60, 160, 220), (90, 200, 120), (240, 200, 80)][k % 4]
    img = Image.new(mode, (size, size + (k % 3) * 7), color)
    img = _draw_variant(img, k, marker)
    if ext in ("jpg", "jpeg"):
        img.convert("RGB").save(path, quality=88)
    elif ext == "ico":
        img.convert("RGBA").save(path, sizes=[(64, 64), (32, 32), (16, 16)])
    else:
        img.save(path)


SVG_TEMPLATE = """<svg xmlns="http://www.w3.org/2000/svg" width="160" height="120">
<rect width="100%" height="100%" fill="#eef"/>
<circle cx="60" cy="50" r="36" fill="#{color}"/>
<path d="M10 110 L80 70 L150 110 Z" fill="#369"/>
<text x="8" y="20" font-size="14" fill="#123">{marker}</text>
</svg>"""


def gen_svg(path: Path, k: int, marker: str) -> None:
    path.write_text(SVG_TEMPLATE.format(color=f"{0x334455 + k * 0x111111:06x}"[-6:],
                                        marker=marker), encoding="utf-8")


def gen_txt(path: Path, k: int, marker: str) -> None:
    paras = [f"这是第{k}号文本变体。包含中文、English、标点！",
             f"标记码 {marker} 用于校验内容完整性。",
             "长段落：" + "内容重复测试。" * (10 + k * 5),
             "", "末段，特殊字符 <>&\"'" + ("🎉" if k % 2 else "★")]
    path.write_text("\n\n".join(paras), encoding="utf-8")


def gen_md(path: Path, k: int, marker: str) -> None:
    path.write_text(f"# {k} 号文档\n\n正文含 **加粗** 与 {marker}。\n\n"
                    f"## 小节\n\n- 列表项一\n- 列表项二{k}\n\n"
                    f"```python\nprint('code {marker}')\n```\n", encoding="utf-8")


def gen_html(path: Path, k: int, marker: str) -> None:
    path.write_text(f"<html><head><meta charset='utf-8'><title>变体{k}</title></head>"
                    f"<body><h1>标题{k}</h1><p>段落内容 <b>加粗</b> {marker}。</p>"
                    f"<p style='color:#cc0000'>红色提示文字。</p>"
                    f"<table border='1'><tr><th>列A</th><th>列B</th></tr>"
                    f"<tr><td>{marker}</td><td>{k}</td></tr></table></body></html>",
                    encoding="utf-8")


def gen_docx(path: Path, k: int, marker: str) -> None:
    from docx import Document
    from docx.shared import RGBColor
    d = Document()
    d.add_heading(f"压测文档 {k}", level=1)
    d.add_paragraph(f"正文段落，标记 {marker}。包含中文与 English mixed text。")
    p = d.add_paragraph()
    r = p.add_run("加粗文字"); r.bold = True
    r = p.add_run("红色文字"); r.font.color.rgb = RGBColor(0xCC, 0, 0)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "列A"; t.cell(0, 1).text = "列B"
    t.cell(1, 0).text = marker; t.cell(1, 1).text = str(k)
    d.save(path)


def gen_pdf(path: Path, k: int, marker: str) -> None:
    import pymupdf
    doc = pymupdf.open()
    for page_no in range(1 + k % 2):
        page = doc.new_page(width=420, height=300)
        page.insert_text((40, 60), f"Stress PDF variant {k}", fontsize=13)
        page.insert_text((40, 90), marker, fontsize=11)
        page.insert_text((40, 120), "The quick brown fox. " * (3 + k), fontsize=9)
    doc.save(path)
    doc.close()


def gen_rtf(path: Path, k: int, marker: str) -> None:
    path.write_text(
        r"{\rtf1\ansi\deff0{\fonttbl{\f0 SimSun;}}"
        f"\n\\fs24 压测RTF变体{k} \\b 加粗 \\b0 {marker}\\par\n第二行内容。\\par\n}}",
        encoding="utf-8")


def _records(k: int, marker: str) -> list[dict]:
    n = [3, 8, 15, 30][k % 4]
    return [{"名称": f"条目{i}", "数量": i * (k + 1), "标记": marker,
             "备注": "中文备注" if i % 2 else "note"} for i in range(1, n + 1)]


def gen_csv_tsv(path: Path, ext: str, k: int, marker: str) -> None:
    import csv
    delim = "\t" if ext == "tsv" else ","
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=delim)
        w.writerow(["名称", "数量", "标记", "备注"])
        for r in _records(k, marker):
            w.writerow([r["名称"], r["数量"], r["标记"], r["备注"]])


def gen_json(path: Path, k: int, marker: str) -> None:
    if k % 3 == 2:  # 嵌套结构变体
        data = {"项目": f"压测{k}", "标记": marker,
                "配置": {"深度": k, "开关": True},
                "列表": [1, 2, 3]}
    else:
        data = _records(k, marker)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def gen_yaml(path: Path, k: int, marker: str) -> None:
    import yaml
    data = _records(k, marker) if k % 2 == 0 else {"名称": f"压测{k}", "标记": marker,
                                                   "嵌套": {"a": 1, "b": [1, 2]}}
    path.write_text(yaml.safe_dump(data, allow_unicode=True, sort_keys=False),
                    encoding="utf-8")


def gen_xml(path: Path, k: int, marker: str) -> None:
    import xmltodict
    payload = {"records": {"record": _records(k, marker)}}
    path.write_text(xmltodict.unparse(payload, pretty=True), encoding="utf-8")


def gen_toml(path: Path, k: int, marker: str) -> None:
    path.write_text(f"title = '压测{k}'\nmarker = '{marker}'\n"
                    f"count = {k + 1}\n[settings]\nenabled = true\nratio = 0.5\n",
                    encoding="utf-8")


def gen_xlsx(path: Path, k: int, marker: str) -> None:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["名称", "数量", "标记", "备注"])
    for r in _records(k, marker):
        ws.append([r["名称"], r["数量"], r["标记"], r["备注"]])
    wb.save(path)


def _ffmpeg(*args: str) -> None:
    ff = media.ffmpeg_path()
    if not ff:
        raise RuntimeError("ffmpeg 不可用")
    proc = subprocess.run([ff, "-y", "-loglevel", "error", *args],
                          creationflags=0x08000000, timeout=120)
    if proc.returncode != 0:
        raise RuntimeError(f"ffmpeg 生成失败: {' '.join(args[:6])}...")


def gen_wav(path: Path, k: int, marker: str) -> None:
    import math
    import struct
    import wave
    freq = [440, 660, 880][k % 3]
    with wave.open(str(path), "w") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(8000)
        frames = bytearray()
        for i in range(8000):  # 1 秒
            sample = int(9000 * math.sin(2 * math.pi * freq * i / 8000))
            frames += struct.pack("<h", sample)
        w.writeframes(bytes(frames))


AUDIO_TARGETS = ["mp3", "flac", "m4a", "aac", "ogg", "opus", "wma"]
VIDEO_SOURCES = ["mp4", "mkv", "avi", "mov", "webm", "wmv"]


def gen_media_set(base_dir: Path) -> list[Path]:
    """生成音频与视频源文件（ffmpeg 合成，内容为音调与测试图案）。"""
    files: list[Path] = []
    wav0 = base_dir / "audio_src0.wav"
    gen_wav(wav0, 0, "AUDIO")
    files.append(wav0)
    for k in (1, 2):
        p = base_dir / f"audio_src{k}.wav"
        gen_wav(p, k, "AUDIO")
        files.append(p)
    for ext in AUDIO_TARGETS:
        p = base_dir / f"audio_src.{ext}"
        _ffmpeg("-i", str(wav0), str(p))
        files.append(p)
    # 视频：测试图案 + 正弦音
    for ext in VIDEO_SOURCES:
        p = base_dir / f"video_src.{ext}"
        codec = {"mp4": ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac"],
                 "mov": ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac"],
                 "mkv": ["-c:v", "libx264", "-c:a", "aac"],
                 "avi": ["-c:v", "mpeg4", "-c:a", "libmp3lame"],
                 "webm": ["-c:v", "libvpx-vp9", "-b:v", "0", "-c:a", "libopus"],
                 "wmv": ["-c:v", "wmv2", "-c:a", "wmav2"]}[ext]
        _ffmpeg("-f", "lavfi", "-i", "testsrc=size=320x240:rate=15:duration=1",
                "-f", "lavfi", "-i", "sine=frequency=440:duration=1",
                *codec, str(p))
        files.append(p)
    return files


# ================================================================ 生成约 200 个文件

def generate_all(base_dir: Path) -> list[dict]:
    files: list[dict] = []

    def add(path: Path, marker: str, has_text: bool):
        files.append({"path": path, "ext": path.suffix.lower().lstrip("."),
                      "marker": marker, "has_text": has_text})

    raster = ["png", "jpg", "webp", "bmp", "gif", "tif", "ico"]
    for k in range(8):  # 7 格式 × 8 变体 = 56
        for ext in raster:
            p = base_dir / f"img_v{k}.{ext}"
            gen_image(p, ext, k, f"ZM{k}")
            add(p, f"ZM{k}", False)
    for k in range(3):
        p = base_dir / f"img_svg_v{k}.svg"
        gen_svg(p, k, f"ZS{k}")
        add(p, f"ZS{k}", True)  # svg 内容含文字
    # heic 仅当编码器可用
    try:
        p = base_dir / "img_v0.heic"
        gen_image(p, "heic", 0, "ZM0")
        add(p, "ZM0", False)
    except Exception:
        pass

    docs = [("txt", gen_txt), ("md", gen_md), ("html", gen_html),
            ("docx", gen_docx), ("pdf", gen_pdf), ("rtf", gen_rtf)]
    for k in range(8):  # 6 格式 × 8 变体 = 48
        for ext, fn in docs:
            p = base_dir / f"doc_v{k}.{ext}"
            fn(p, k, f"ZD{k}")
            add(p, f"ZD{k}", True)

    data = [("csv", gen_csv_tsv), ("tsv", gen_csv_tsv), ("json", gen_json),
            ("yaml", gen_yaml), ("xml", gen_xml), ("toml", gen_toml), ("xlsx", gen_xlsx)]
    for k in range(8):  # 7 格式 × 8 变体 = 56
        for ext, fn in data:
            p = base_dir / f"data_v{k}.{ext}"
            if ext in ("csv", "tsv"):
                fn(p, ext, k, f"ZL{k}")
            else:
                fn(p, k, f"ZL{k}")
            add(p, f"ZL{k}", True)

    media_dir = base_dir / "media"
    media_dir.mkdir(exist_ok=True)
    for f in gen_media_set(media_dir):  # 3 wav + 7 音频 + 6 视频 = 16
        files.append({"path": f, "ext": f.suffix.lower().lstrip("."),
                      "marker": "AUDIO", "has_text": False})

    # PPTX 源（python-pptx 生成）
    from pptx import Presentation
    for k in range(3):
        p = base_dir / f"ppt_v{k}.pptx"
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"压测幻灯 {k}"
        s.placeholders[1].text = f"幻灯内容标记 ZP{k}"
        prs.save(p)
        files.append({"path": p, "ext": "pptx", "marker": f"ZP{k}", "has_text": True})

    # EPUB 源
    from core.epub_tools import build_epub
    for k in range(2):
        p = base_dir / f"book_v{k}.epub"
        build_epub([(f"第一章", f"<p>Book {k} ZE{k} 段落一。</p><p>段落二。</p>"),
                    (f"第二章", f"<p>第二章内容 ZE{k} 结尾。</p>")], p, f"压测书{k}")
        files.append({"path": p, "ext": "epub", "marker": f"ZE{k}", "has_text": True})

    # 伪装 docx（网页内容冒充）
    for k in range(2):
        p = base_dir / f"fake_v{k}.docx"
        p.write_text(f"<html><body><h1>伪装{k}</h1><p>伪装内容 ZF{k}。</p></body></html>",
                     encoding="utf-8")
        files.append({"path": p, "ext": "docx", "marker": f"ZF{k}", "has_text": True})

    return files


# ================================================================ 组合枚举与执行

def legal_targets(ext: str) -> list[str]:
    targets: list[str] = []
    for _, mod in MODULES:
        for t, _label in mod.TARGETS:
            if t != ext and t not in targets and mod.can_convert(ext, t):
                targets.append(t)
    return targets


MAGIC = {
    "png": (0, b"\x89PNG"), "jpg": (0, b"\xff\xd8"), "webp": (8, b"WEBP"),
    "bmp": (0, b"BM"), "gif": (0, b"GIF8"), "tif": (0, b"II*\x00"),
    "pdf": (0, b"%PDF"), "docx": (0, b"PK\x03\x04"), "xlsx": (0, b"PK\x03\x04"),
    "wav": (0, b"RIFF"), "flac": (0, b"fLaC"), "ogg": (0, b"OggS"),
    "opus": (0, b"OggS"), "wma": (0, b"\x30\x26\xb2\x75"),
    "mp4": (4, b"ftyp"), "mov": (4, b"ftyp"), "mkv": (0, b"\x1a\x45\xdf\xa3"),
    "webm": (0, b"\x1a\x45\xdf\xa3"), "avi": (0, b"RIFF"),
}


def validate(src: dict, dst: str, outputs: list[Path]) -> str | None:
    """校验输出；返回错误描述，None 为通过。"""
    if not outputs:
        return "无输出文件"
    for o in outputs:
        if not o.exists():
            return f"输出不存在: {o.name}"
        if o.stat().st_size <= 0:
            return f"输出为空文件: {o.name}"
    # 魔数抽查（每个任务抽查第一个输出）
    if dst in MAGIC:
        off, sig = MAGIC[dst]
        head = outputs[0].read_bytes()
        if len(head) <= off or head[off:off + len(sig)] != sig:
            return f"魔数不符: 期望 {sig!r}"
    if dst in ("png", "jpg", "webp", "bmp", "gif", "tif", "ico"):
        try:
            with Image.open(outputs[0]) as im:
                im.verify()
        except Exception as e:
            return f"图片损坏: {e}"
    if dst == "pdf" and src["has_text"] and src["ext"] not in ("png", "jpg", "webp",
                                                              "bmp", "gif", "tif", "ico",
                                                              "svg", "heic", "heif"):
        try:
            import pymupdf
            with pymupdf.open(outputs[0]) as doc:
                text = "".join(page.get_text() for page in doc)
            if src["marker"] not in text:
                return f"PDF 内容缺标记 {src['marker']}"
        except ConvertError:
            raise
        except Exception as e:
            return f"PDF 打不开: {e}"
    if dst == "docx" and src["has_text"]:
        try:
            import docx as docx_lib
            d = docx_lib.Document(str(outputs[0]))
            text = "\n".join(p.text for p in d.paragraphs)
            if src["marker"] not in text:
                return f"DOCX 内容缺标记 {src['marker']}"
        except Exception as e:
            return f"DOCX 打不开: {e}"
    if dst in ("txt", "md", "html") and src["has_text"]:
        if src["marker"] not in outputs[0].read_text(encoding="utf-8", errors="replace"):
            return f"内容缺标记 {src['marker']}"
    if dst in ("csv", "tsv", "json", "yaml", "yml", "xml", "toml") and src["has_text"]:
        if src["marker"] not in outputs[0].read_text(encoding="utf-8", errors="replace"):
            return f"内容缺标记 {src['marker']}"
    if dst == "epub" and src["has_text"]:
        import zipfile as _zf
        with _zf.ZipFile(outputs[0]) as z:
            content = b"".join(z.read(n) for n in z.namelist()
                               if n.endswith(".xhtml")).decode("utf-8", "replace")
        if src["marker"] not in content:
            return "EPUB 内容缺标记"
    if dst == "xlsx" and src["has_text"]:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(outputs[0], read_only=True)
            found = any(src["marker"] in str(c)
                        for ws in wb.worksheets for row in ws.iter_rows(values_only=True)
                        for c in row if c is not None)
            wb.close()
            if not found:
                return f"XLSX 内容缺标记 {src['marker']}"
        except Exception as e:
            return f"XLSX 打不开: {e}"
    return None


def main() -> int:
    base = Path(tempfile.mkdtemp(prefix="stress_"))
    files_dir = base / "files"
    files_dir.mkdir()
    print("【1/3】生成源文件…", flush=True)
    files = generate_all(files_dir)
    print(f"  共生成 {len(files)} 个源文件", flush=True)

    # 枚举任务：每个文件 × 每个合法目标
    tasks: list[tuple[dict, str]] = []
    pairs: set[tuple[str, str]] = set()
    for f in files:
        for dst in legal_targets(f["ext"]):
            tasks.append((f, dst))
            pairs.add((f["ext"], dst))
    legal_pairs = {(e, t) for e in {f["ext"] for f in files} for t in legal_targets(e)}
    print(f"【2/3】执行 {len(tasks)} 次转换（{len(pairs)} 种组合，并行 {PARALLEL}）…", flush=True)

    results: list[dict] = []

    def run_one(item: tuple[dict, str]) -> dict:
        src, dst = item
        try:
            res = registry.convert_file(src["path"], dst, base / "out",
                                        {"quality": 95, "dpi": 150, "media_level": "max",
                                         "pdf_password": "压测Pass1"})
            err = validate(src, dst, res.outputs)
            status = "ok" if err is None else "bad"
            return {"src": src, "dst": dst, "status": status, "err": err,
                    "out": [str(o) for o in res.outputs]}
        except ConvertError as e:
            return {"src": src, "dst": dst, "status": "fail", "err": str(e)[:200], "out": []}
        except Exception as e:
            return {"src": src, "dst": dst, "status": "fail",
                    "err": f"{type(e).__name__}: {e}"[:200], "out": []}

    with ThreadPoolExecutor(max_workers=PARALLEL) as pool:
        futs = [pool.submit(run_one, t) for t in tasks]
        done = 0
        for fut in as_completed(futs):
            results.append(fut.result())
            done += 1
            if done % 100 == 0:
                print(f"  进度 {done}/{len(tasks)}", flush=True)

    # 【3/3】汇总
    ok = [r for r in results if r["status"] == "ok"]
    bad = [r for r in results if r["status"] == "bad"]
    fail = [r for r in results if r["status"] == "fail"]
    tested_pairs = {(r["src"]["ext"], r["dst"]) for r in results}
    missing = legal_pairs - tested_pairs

    lines = [f"压测报告  {base}",
             f"源文件: {len(files)} 个 | 转换执行: {len(results)} 次",
             f"通过: {len(ok)} | 校验失败: {len(bad)} | 转换报错: {len(fail)}",
             f"组合覆盖: {len(tested_pairs)}/{len(legal_pairs)} 种"
             + ("" if not missing else f"  缺失: {sorted(missing)}")]
    if bad or fail:
        lines.append("---- 失败清单 ----")
        for r in (fail + bad):
            lines.append(f"  ✗ {r['src']['path'].name} → {r['dst']}: {r['err']}")
    report = "\n".join(lines)
    (ROOT / "tests" / "stress_report.txt").write_text(report, encoding="utf-8")
    print("\n" + report)
    return 0 if not bad and not fail and not missing else 1


if __name__ == "__main__":
    sys.exit(main())
